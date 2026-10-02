#!/usr/bin/env python3
"""
OASIS v2.0 — Carga desde el Sheets exportado hacia Cloud SQL.

Se puede correr tantas veces como haga falta: hace upsert por llave natural,
así que refrescar los datos mientras los equipos siguen llenando 2026 es
simplemente volver a ejecutarlo.

Uso:
    source .env
    python cargar_oasis.py --sheet 13Vh45CymNMntNG-EW3NQnglnfyvq39nT7nvediDTV1g   # lee el Sheets en vivo
    python cargar_oasis.py oasis.xlsx                                           # o un Excel descargado

Opciones:
    --sheet ID        ID del Google Sheets (el de la URL). Requiere credenciales_google.json
    --credenciales    ruta a la llave de la cuenta de servicio (default: credenciales_google.json)
    --mercado MCO     mercado al que pertenece el archivo (default: MCO)
    --dry-run         muestra qué haría sin escribir nada
"""

import argparse
import os
import re
import sys
from collections import defaultdict

import openpyxl
import psycopg2
import psycopg2.extras


# ---------------------------------------------------------------------------
# Localización de encabezados.
#
# El script NO usa números de fila ni de columna fijos. Busca la fila de
# encabezados por un texto ancla y mapea cada columna por su nombre. Si el
# Sheets cambia de layout, se adapta; si falta un encabezado requerido, se
# detiene y dice cuál.
# ---------------------------------------------------------------------------

import unicodedata


def normalizar(v):
    """'  Sub Campaña\n(x) ' -> 'SUB CAMPANA (X)'. Sin acentos, sin saltos."""
    if v is None:
        return ""
    s = unicodedata.normalize("NFD", str(v))
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return " ".join(s.split()).upper()


def localizar(ws, ancla_col, ancla_texto, columnas, hoja):
    """
    Devuelve (fila_datos, {clave: indice_columna}).

    ancla_col    : columna (1-based) donde debe aparecer ancla_texto
    ancla_texto  : texto normalizado que identifica la fila de encabezados
    columnas     : {clave: encabezado_esperado_normalizado}
    """
    fila_hdr = None
    for r in range(1, min(ws.max_row, 50) + 1):
        if normalizar(ws.cell(row=r, column=ancla_col).value) == ancla_texto:
            fila_hdr = r
            break
    if fila_hdr is None:
        sys.exit(f"[{hoja}] No encontré la fila de encabezados: "
                 f"esperaba '{ancla_texto}' en la columna {ancla_col}.")

    encabezados = {}
    for c in range(1, ws.max_column + 1):
        h = normalizar(ws.cell(row=fila_hdr, column=c).value)
        if h and h not in encabezados:
            encabezados[h] = c

    mapa, faltan = {}, []
    for clave, esperado in columnas.items():
        idx = encabezados.get(esperado)
        if idx is None:
            # tolera encabezados con texto adicional, ej. "MATT CAMPAIGN TYPE (NO ESCRIBIR A MANO)"
            idx = next((i for h, i in encabezados.items() if h.startswith(esperado)), None)
        if idx is None:
            faltan.append(f"{clave} ('{esperado}')")
        else:
            mapa[clave] = idx

    if faltan:
        sys.exit(f"[{hoja}] Faltan encabezados en la fila {fila_hdr}: " + ", ".join(faltan)
                 + f"\n  Encabezados encontrados: {sorted(encabezados)}")

    return fila_hdr + 1, mapa


FLOW_COLS = {
    "codigo": "ID", "cliente": "CLIENTE", "campana": "CAMPANA",
    "tipo_compra": "TIPO DE COMPRA", "proveedor": "PROVEEDOR", "medio": "MEDIO",
    "tipo_costo": "TIPO DE COSTO", "formato": "FORMATO", "ubicacion": "UBICACION",
    "trafico": "TRAFICO", "ciudad": "CIUDAD", "tiempo": "TIEMPO",
    "tarifa_bruta": "TARIFA BRUTA", "descuento_pct": "DESC %", "tarifa_neta": "TARIFA NETA",
    "cantidad": "CANTIDAD ELEMENTOS", "nro_semanas": "NRO. SEMANAS",
    "valor_total": "VALOR TOTAL", "fecha_inicio": "FECHA INICIO", "fecha_fin": "FECHA FIN",
}

TRACKING_COLS = {
    "material": "ID FLOW", "ejecucion": "ID EJECUCION", "sub_campana": "SUB CAMPANA",
    "referencia": "REFERENCIA", "enlace": "ENLACE MATERIAL",
    "reporte": "REPORTE DE IMPLEMENTACION",
    "fecha_inicio": "FECHA INICIO", "fecha_fin": "FECHA FIN",
}

GLOSARIO_COLS = {
    "nombre_calendario": "NOMBRE DE CAMPANA EN CALENDARIO", "mercado": "MERCADO",
    "categorizacion": "CATEGORIZACION", "nombre_unico": "NOMBRE UNICO DE CAMPANA",
    "matt_campaign_type": "MATT CAMPAIGN TYPE",
}


# ---------------------------------------------------------------------------
# Lectura directa de Google Sheets.
#
# Presenta las hojas con la misma interfaz que openpyxl (.cell(row, col).value,
# .max_row, .max_column) para que el resto del script no distinga el origen.
# ---------------------------------------------------------------------------

from datetime import date, timedelta

_EPOCA_SHEETS = date(1899, 12, 30)   # el serial 1 de Sheets es el 31/12/1899


class _Celda:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class _HojaSheets:
    def __init__(self, filas):
        self._filas = filas
        self.max_row = len(filas)
        self.max_column = max((len(f) for f in filas), default=0)

    def cell(self, row, column):
        try:
            v = self._filas[row - 1][column - 1]
        except IndexError:
            return _Celda(None)
        return _Celda(None if v == "" else v)


class _LibroSheets:
    def __init__(self, sheet_id, credenciales):
        try:
            import gspread
        except ImportError:
            sys.exit("Falta la librería gspread. Instálala con: pip3 install gspread")
        if not os.path.exists(credenciales):
            sys.exit(f"No encuentro {credenciales}. Es la llave JSON de la cuenta de servicio.")
        cliente = gspread.service_account(filename=credenciales)
        try:
            self._libro = cliente.open_by_key(sheet_id)
        except gspread.exceptions.APIError as e:
            sys.exit(f"No pude abrir el Sheets. ¿Lo compartiste con la cuenta de servicio?\n{e}")
        self._cache = {}

    def __getitem__(self, nombre):
        if nombre not in self._cache:
            try:
                hoja = self._libro.worksheet(nombre)
            except Exception:
                sys.exit(f"El Sheets no tiene una hoja llamada '{nombre}'.")
            filas = hoja.get_all_values(value_render_option="UNFORMATTED_VALUE")
            self._cache[nombre] = _HojaSheets(filas)
        return self._cache[nombre]


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def texto(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def numero(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^0-9.\-]", "", str(v))
    try:
        return float(s) if s not in ("", "-", ".") else None
    except ValueError:
        return None


def fecha(v):
    """Acepta datetime (Excel) o serial numérico (Sheets). Texto -> None."""
    if v is None or v == "":
        return None
    if hasattr(v, "date"):
        return v.date()
    if isinstance(v, (int, float)) and 20000 <= v <= 80000:   # ~1954 a ~2119
        return _EPOCA_SHEETS + timedelta(days=int(v))
    return None


def extraer_campana(taxonomia):
    """
    'MCO_ML_OTH-PF_ALL_PARQUE-FIJO_X' -> ['PARQUE-FIJO_X', 'PARQUE-FIJO']

    Devuelve candidatos en orden de preferencia: el extracto completo después
    del cuarto guion bajo, y el mismo sin un sufijo de una sola letra.
    """
    if not taxonomia:
        return []
    m = re.match(r"^(?:[^_]*_){4}(.*)$", taxonomia)
    if not m:
        return [taxonomia]
    bruto = m.group(1)
    candidatos = [bruto]
    sin_sufijo = re.sub(r"_[A-Z]$", "", bruto)
    if sin_sufijo != bruto:
        candidatos.append(sin_sufijo)
    return candidatos


def secuencia_desde_codigo(codigo):
    """'001-C' -> 3. Devuelve None si el formato no calza."""
    m = re.match(r"^(.+)-([A-Z])$", str(codigo).strip())
    if not m:
        return None, None
    return m.group(1), ord(m.group(2)) - 64


def filas(ws, desde, col_clave):
    for r in range(desde, ws.max_row + 1):
        if texto(ws.cell(row=r, column=col_clave).value):
            yield r


# ---------------------------------------------------------------------------
# Carga
# ---------------------------------------------------------------------------

class Carga:
    def __init__(self, conn, wb, mercado_codigo, dry_run=False):
        self.cur = conn.cursor()
        self.conn = conn
        self.wb = wb
        self.mercado_codigo = mercado_codigo
        self.dry_run = dry_run
        self.avisos = defaultdict(list)
        self.stats = defaultdict(int)

        self.flow_ini, self.FLOW = localizar(wb["Flow"], 1, "ID", FLOW_COLS, "Flow")
        self.trk_ini,  self.TRK  = localizar(wb["Tracking"], 1, "ID FLOW", TRACKING_COLS, "Tracking")
        self.glo_ini,  self.GLO  = localizar(wb["Glosario"], 3, "MERCADO", GLOSARIO_COLS, "Glosario")
        print(f"  Flow: datos desde fila {self.flow_ini} | Tracking: desde fila {self.trk_ini} | Glosario: desde fila {self.glo_ini}")

        self.cur.execute("SELECT id FROM mercados WHERE codigo = %s", (mercado_codigo,))
        fila = self.cur.fetchone()
        if not fila:
            sys.exit(f"El mercado {mercado_codigo} no existe en la tabla mercados.")
        self.mercado_id = fila[0]

    # -- campañas ----------------------------------------------------------

    def campanas(self):
        ws = self.wb["Glosario"]
        datos = []
        for r in filas(ws, self.glo_ini, self.GLO["nombre_unico"]):
            mercado = texto(ws.cell(row=r, column=self.GLO["mercado"]).value)
            nombre = texto(ws.cell(row=r, column=self.GLO["nombre_unico"]).value)
            if not mercado or not nombre:
                continue
            datos.append((
                mercado.upper(), nombre,
                texto(ws.cell(row=r, column=self.GLO["nombre_calendario"]).value),
                texto(ws.cell(row=r, column=self.GLO["categorizacion"]).value),
                texto(ws.cell(row=r, column=self.GLO["matt_campaign_type"]).value),
            ))

        vistos = set()
        unicos = []
        for d in datos:
            if (d[0], d[1]) in vistos:
                continue
            vistos.add((d[0], d[1]))
            unicos.append(d)

        psycopg2.extras.execute_batch(self.cur, """
            INSERT INTO campanas
                (mercado_id, nombre_unico, nombre_calendario,
                 categorizacion, matt_campaign_type, sincronizado_en)
            SELECT m.id, %s, %s, %s, %s, now()
            FROM mercados m WHERE m.codigo = %s
            ON CONFLICT (mercado_id, nombre_unico) DO UPDATE SET
                nombre_calendario  = EXCLUDED.nombre_calendario,
                categorizacion     = EXCLUDED.categorizacion,
                matt_campaign_type = EXCLUDED.matt_campaign_type,
                sincronizado_en    = now()
        """, [(d[1], d[2], d[3], d[4], d[0]) for d in unicos])

        self.stats["campanas"] = len(unicos)

        self.cur.execute("""
            SELECT c.nombre_unico, c.id FROM campanas c WHERE c.mercado_id = %s
        """, (self.mercado_id,))
        self.mapa_campanas = dict(self.cur.fetchall())

    # -- proveedores -------------------------------------------------------

    def proveedores(self):
        """El catálogo no existe hoy: se deriva de los valores del Flow."""
        ws = self.wb["Flow"]
        nombres = set()
        for r in filas(ws, self.flow_ini, self.FLOW["codigo"]):
            n = texto(ws.cell(row=r, column=self.FLOW["proveedor"]).value)
            if n:
                nombres.add(n.upper())

        psycopg2.extras.execute_batch(self.cur, """
            INSERT INTO proveedores (mercado_id, nombre)
            VALUES (%s, %s)
            ON CONFLICT (mercado_id, nombre) DO NOTHING
        """, [(self.mercado_id, n) for n in sorted(nombres)])

        self.stats["proveedores"] = len(nombres)

        self.cur.execute(
            "SELECT nombre, id FROM proveedores WHERE mercado_id = %s",
            (self.mercado_id,))
        self.mapa_proveedores = dict(self.cur.fetchall())

    # -- materiales --------------------------------------------------------

    def materiales(self):
        ws = self.wb["Flow"]
        pendientes = []

        for r in filas(ws, self.flow_ini, self.FLOW["codigo"]):
            g = lambda k: ws.cell(row=r, column=self.FLOW[k]).value
            codigo = texto(g("codigo"))

            fi, ff = fecha(g("fecha_inicio")), fecha(g("fecha_fin"))
            if not fi or not ff:
                self.avisos["material sin fechas"].append(codigo)
                continue
            if ff < fi:
                self.avisos["material con fechas invertidas"].append(codigo)
                continue

            campana_id = None
            for cand in extraer_campana(texto(g("campana"))):
                if cand in self.mapa_campanas:
                    campana_id = self.mapa_campanas[cand]
                    break
            if not campana_id:
                self.avisos["campaña no encontrada en glosario"].append(
                    f"{codigo}: {texto(g('campana'))}")
                continue

            prov = (texto(g("proveedor")) or "").upper()
            proveedor_id = self.mapa_proveedores.get(prov)
            if not proveedor_id:
                self.avisos["proveedor no encontrado"].append(f"{codigo}: {prov}")
                continue

            pendientes.append((
                self.mercado_id, codigo, texto(g("cliente")),
                campana_id, proveedor_id,
                texto(g("tipo_compra")), texto(g("medio")), texto(g("tipo_costo")),
                texto(g("formato")), texto(g("ubicacion")), texto(g("ciudad")),
                numero(g("trafico")), texto(g("tiempo")),
                numero(g("tarifa_bruta")) or 0, numero(g("descuento_pct")) or 0,
                numero(g("tarifa_neta")) or 0, int(numero(g("cantidad")) or 1),
                numero(g("nro_semanas")), numero(g("valor_total")) or 0,
                fi, ff,
            ))

        psycopg2.extras.execute_batch(self.cur, """
            INSERT INTO materiales (
                mercado_id, codigo, cliente, campana_id, proveedor_id,
                tipo_compra, medio, tipo_costo, formato, ubicacion, ciudad,
                trafico, tiempo, tarifa_bruta, descuento_pct, tarifa_neta,
                cantidad, nro_semanas, valor_total, fecha_inicio, fecha_fin)
            VALUES (%s,%s,%s,%s,%s,%s::tipo_compra_enum,%s::medio_enum,
                    %s::tipo_costo_enum,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            ON CONFLICT (mercado_id, codigo) DO UPDATE SET
                cliente       = EXCLUDED.cliente,
                campana_id    = EXCLUDED.campana_id,
                proveedor_id  = EXCLUDED.proveedor_id,
                tipo_compra   = EXCLUDED.tipo_compra,
                medio         = EXCLUDED.medio,
                tipo_costo    = EXCLUDED.tipo_costo,
                formato       = EXCLUDED.formato,
                ubicacion     = EXCLUDED.ubicacion,
                ciudad        = EXCLUDED.ciudad,
                trafico       = EXCLUDED.trafico,
                tiempo        = EXCLUDED.tiempo,
                tarifa_bruta  = EXCLUDED.tarifa_bruta,
                descuento_pct = EXCLUDED.descuento_pct,
                tarifa_neta   = EXCLUDED.tarifa_neta,
                cantidad      = EXCLUDED.cantidad,
                nro_semanas   = EXCLUDED.nro_semanas,
                valor_total   = EXCLUDED.valor_total,
                fecha_inicio  = EXCLUDED.fecha_inicio,
                fecha_fin     = EXCLUDED.fecha_fin,
                actualizado_en = now()
        """, pendientes)

        self.stats["materiales"] = len(pendientes)

        self.cur.execute(
            "SELECT codigo, id FROM materiales WHERE mercado_id = %s",
            (self.mercado_id,))
        self.mapa_materiales = dict(self.cur.fetchall())

    # -- ejecuciones -------------------------------------------------------

    def ejecuciones(self):
        ws = self.wb["Tracking"]
        pendientes = []

        for r in filas(ws, self.trk_ini, self.TRK["ejecucion"]):
            g = lambda k: ws.cell(row=r, column=self.TRK[k]).value
            codigo = texto(g("ejecucion"))

            mat_codigo, secuencia = secuencia_desde_codigo(codigo)
            if secuencia is None:
                self.avisos["código de ejecución con formato inesperado"].append(codigo)
                continue

            material_id = self.mapa_materiales.get(mat_codigo)
            if not material_id:
                self.avisos["ejecución sin material cargado"].append(codigo)
                continue

            # La sub campaña puede traer valores fuera del glosario (ODM,
            # espacios al final). No bloquean la carga: entran en nulo y
            # quedan listadas en v_alertas.
            sub = texto(g("sub_campana"))
            sub_id = self.mapa_campanas.get(sub) if sub else None
            if sub and not sub_id:
                self.avisos["sub campaña fuera del glosario"].append(
                    f"{codigo}: {sub!r}")

            raw_fi, raw_ff = g("fecha_inicio"), g("fecha_fin")
            fi, ff = fecha(raw_fi), fecha(raw_ff)
            if (fi is None) != (ff is None):
                self.avisos["ejecución con una sola fecha"].append(codigo)
                fi = ff = None
            elif fi is None:
                # Distingue celda vacía de texto que parece fecha pero no lo es
                if texto(raw_fi) or texto(raw_ff):
                    self.avisos["ejecución con fecha escrita como texto (no es fecha válida)"].append(
                        f"{codigo}: {texto(raw_fi)!r} → {texto(raw_ff)!r}")
                else:
                    self.avisos["ejecución sin fechas"].append(codigo)
            if fi and ff and ff < fi:
                self.avisos["ejecución con fechas invertidas"].append(codigo)
                continue

            pendientes.append((
                material_id, secuencia, sub_id,
                texto(g("referencia")), texto(g("enlace")), texto(g("reporte")),
                fi, ff,
            ))

        psycopg2.extras.execute_batch(self.cur, """
            INSERT INTO ejecuciones (
                material_id, secuencia, sub_campana_id,
                referencia, enlace, reporte_implementacion,
                fecha_inicio, fecha_fin)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
            ON CONFLICT (material_id, secuencia) DO UPDATE SET
                sub_campana_id         = EXCLUDED.sub_campana_id,
                referencia             = EXCLUDED.referencia,
                enlace                 = EXCLUDED.enlace,
                reporte_implementacion = EXCLUDED.reporte_implementacion,
                fecha_inicio           = EXCLUDED.fecha_inicio,
                fecha_fin              = EXCLUDED.fecha_fin,
                actualizado_en         = now()
        """, pendientes)

        self.stats["ejecuciones"] = len(pendientes)

    # -- huérfanos ---------------------------------------------------------

    def huerfanos(self):
        """Filas que existen en la base pero ya no en el Sheets."""
        ws = self.wb["Flow"]
        codigos = {texto(ws.cell(row=r, column=self.FLOW["codigo"]).value)
                   for r in filas(ws, self.flow_ini, self.FLOW["codigo"])}
        self.cur.execute(
            "SELECT codigo FROM materiales WHERE mercado_id = %s", (self.mercado_id,))
        for (c,) in self.cur.fetchall():
            if c not in codigos:
                self.avisos["material en la base que ya no está en el Sheets"].append(c)

    # -- reporte -----------------------------------------------------------

    def reporte(self):
        print("\n" + "=" * 62)
        print(f"  Carga de {self.mercado_codigo}" + ("  (simulación)" if self.dry_run else ""))
        print("=" * 62)
        for k in ("campanas", "proveedores", "materiales", "ejecuciones"):
            print(f"  {k:<14} {self.stats[k]:>6}")

        if self.avisos:
            print("\n  Avisos")
            for motivo, items in sorted(self.avisos.items()):
                print(f"\n  {motivo} ({len(items)})")
                for i in items[:10]:
                    print(f"      {i}")
                if len(items) > 10:
                    print(f"      ... y {len(items) - 10} más")
        else:
            print("\n  Sin avisos.")

        print("\n  Verificación: SELECT * FROM v_cuadre;  (debe salir vacío)")
        print("  Pendientes:   SELECT * FROM v_alertas;")
        print()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("archivo", nargs="?", help="Excel exportado del Sheets")
    ap.add_argument("--sheet", help="ID del Google Sheets para leer en vivo")
    ap.add_argument("--credenciales", default="credenciales_google.json")
    ap.add_argument("--mercado", default="MCO")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not args.sheet and not args.archivo:
        ap.error("indica un archivo .xlsx o --sheet ID")

    if not os.getenv("PGHOST"):
        sys.exit("Faltan las variables de conexión. Corre: source .env")

    if args.sheet:
        print(f"  Leyendo Google Sheets {args.sheet} en vivo...")
        wb = _LibroSheets(args.sheet, args.credenciales)
    else:
        wb = openpyxl.load_workbook(args.archivo, data_only=True)
    conn = psycopg2.connect()

    try:
        carga = Carga(conn, wb, args.mercado.upper(), args.dry_run)
        carga.campanas()
        carga.proveedores()
        carga.materiales()
        carga.ejecuciones()
        carga.huerfanos()

        if args.dry_run:
            conn.rollback()
        else:
            conn.commit()

        carga.reporte()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    main()
